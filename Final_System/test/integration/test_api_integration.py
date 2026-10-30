"""
End-to-end API tests against a real SQL Server TEST database (marker `db`).
Flow: accounts/auth -> areas/cameras -> Desktop ingestion -> lookup/processing -> reports -> retention.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timedelta

import cv2
import numpy as np
import pytest

pytestmark = pytest.mark.db

ADMIN_PASSWORD = "Admin12345"
CTSV_PASSWORD = "Ctsv12345"


def jpeg(width=80, height=160, seed=3) -> bytes:
    image = np.random.default_rng(seed).integers(0, 255, size=(height, width, 3), dtype=np.uint8)
    return cv2.imencode(".jpg", image)[1].tobytes()


def local_now(client) -> datetime:
    return client.app.state.settings.now_local()


@pytest.fixture(scope="module")
def admin(client, database):
    from api.repositories import accounts

    accounts.create_account(database, "pytest_admin", "Quản trị Pytest", "ADMIN", ADMIN_PASSWORD)
    response = client.post("/api/v1/auth/login", json={"username": "pytest_admin", "password": ADMIN_PASSWORD})
    assert response.status_code == 200, response.text
    token = response.json()["data"]["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def ctsv(client, admin):
    response = client.post("/api/v1/accounts", headers=admin, json={
        "username": "pytest_ctsv", "full_name": "Nhân viên CTSV", "role": "CTSV", "password": CTSV_PASSWORD})
    assert response.status_code == 201, response.text
    login = client.post("/api/v1/auth/login", json={"username": "pytest_ctsv", "password": CTSV_PASSWORD})
    return {"Authorization": f"Bearer {login.json()['data']['access_token']}"}


@pytest.fixture(scope="module")
def setup_camera(client, admin):
    area = client.post("/api/v1/areas", headers=admin, json={"code": "cong-a", "name": "Cổng A", "description": "Cổng chính"})
    assert area.status_code == 201, area.text
    area2 = client.post("/api/v1/areas", headers=admin, json={"code": "THU_VIEN", "name": "Thư viện"})
    camera = client.post("/api/v1/cameras", headers=admin, json={
        "code": "CAM_A1", "name": "Camera cổng A", "area_id": area.json()["data"]["id"],
        "location": "Trụ cổng trái", "stream_source": "rtsp://user:pass@10.0.0.5/stream"})
    assert camera.status_code == 201, camera.text
    camera2 = client.post("/api/v1/cameras", headers=admin, json={
        "code": "CAM_TV1", "name": "Camera thư viện", "area_id": area2.json()["data"]["id"]})
    body, body2 = camera.json(), camera2.json()
    return {
        "area_id": area.json()["data"]["id"], "area2_id": area2.json()["data"]["id"],
        "camera_id": body["data"]["id"], "camera2_id": body2["data"]["id"],
        "device": {"X-Camera-Code": "CAM_A1", "X-Device-Key": body["device_key"]},
        "device2": {"X-Camera-Code": "CAM_TV1", "X-Device-Key": body2["device_key"]},
    }


def send(client, headers, code="UNIFORM", occurred_at=None, event_id=None, image=None, extra=None):
    data = {"event_id": event_id or str(uuid.uuid4()), "violation_code": code,
            "occurred_at": (occurred_at or local_now(client)).isoformat(), "track_id": "7",
            "uniform_confidence": "0.91", "card_confidence": "0.87"}
    data.update(extra or {})
    files = {"image": ("crop.jpg", io.BytesIO(jpeg() if image is None else image), "image/jpeg")}
    return client.post("/api/v1/violations", headers=headers, data=data, files=files)


# ------------------------------------------------------------------------------- health/auth

def test_health_reports_database_without_secrets(client):
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    body = response.json()
    assert body["database"]["status"] == "ok" and {"001", "007", "008"} <= set(body["database"]["migrations"])
    text = response.text.lower()
    assert "pwd" not in text and "password" not in text and "driver=" not in text and "sqlexpress" not in text
    assert client.get("/health").json() == {"status": "ok"}


def test_login_errors_and_lockout(client, database, admin):
    from api.repositories import accounts

    accounts.create_account(database, "pytest_lock", "Khóa thử", "CTSV", "Lock12345")
    for _ in range(3):
        wrong = client.post("/api/v1/auth/login", json={"username": "pytest_lock", "password": "Sai12345"})
        assert wrong.status_code == 401 and wrong.json()["error"]["code"] == "invalid_credentials"
    locked = client.post("/api/v1/auth/login", json={"username": "pytest_lock", "password": "Lock12345"})
    assert locked.status_code == 423
    unknown = client.post("/api/v1/auth/login", json={"username": "khong_ton_tai", "password": "Abcdef12"})
    assert unknown.status_code == 401
    assert client.get("/api/v1/auth/me").status_code == 401
    assert client.get("/api/v1/auth/me", headers={"Authorization": "Bearer abc.def"}).status_code == 401


def test_me_and_role_based_access(client, admin, ctsv):
    assert client.get("/api/v1/auth/me", headers=admin).json()["data"]["role"] == "ADMIN"
    assert client.get("/api/v1/auth/me", headers=ctsv).json()["data"]["role"] == "CTSV"
    assert client.get("/api/v1/accounts", headers=ctsv).status_code == 403
    assert client.post("/api/v1/areas", headers=ctsv, json={"code": "X", "name": "X"}).status_code == 403
    accounts_list = client.get("/api/v1/accounts", headers=admin).json()["items"]
    assert all("PASSWORD_HASH" not in a and "password" not in a for a in accounts_list)


def test_password_policy_and_change_password_revokes_token(client, admin):
    weak = client.post("/api/v1/accounts", headers=admin, json={
        "username": "pytest_weak", "full_name": "Yếu", "role": "CTSV", "password": "123"})
    assert weak.status_code == 422
    created = client.post("/api/v1/accounts", headers=admin, json={
        "username": "pytest_pw", "full_name": "Đổi mật khẩu", "role": "CTSV", "password": "Start12345"})
    assert created.status_code == 201
    token = client.post("/api/v1/auth/login", json={"username": "pytest_pw", "password": "Start12345"}).json()
    headers = {"Authorization": f"Bearer {token['data']['access_token']}"}
    bad = client.post("/api/v1/auth/change-password", headers=headers,
                      json={"current_password": "Wrong12345", "new_password": "Next12345"})
    assert bad.status_code == 400
    ok = client.post("/api/v1/auth/change-password", headers=headers,
                     json={"current_password": "Start12345", "new_password": "Next12345"})
    assert ok.status_code == 200
    assert client.get("/api/v1/auth/me", headers=headers).status_code == 401  # old token revoked
    assert client.post("/api/v1/auth/login", json={"username": "pytest_pw", "password": "Next12345"}).status_code == 200


def test_admin_protection_rules(client, admin):
    me = client.get("/api/v1/auth/me", headers=admin).json()["data"]
    self_disable = client.put(f"/api/v1/accounts/{me['id']}", headers=admin,
                              json={"full_name": "Quản trị Pytest", "role": "ADMIN", "active": False})
    assert self_disable.status_code == 409
    duplicate = client.post("/api/v1/accounts", headers=admin, json={
        "username": "pytest_admin", "full_name": "Trùng", "role": "ADMIN", "password": "Abcdef123"})
    assert duplicate.status_code == 409


def test_disabled_account_token_is_revoked(client, admin):
    created = client.post("/api/v1/accounts", headers=admin, json={
        "username": "pytest_off", "full_name": "Sắp khóa", "role": "CTSV", "password": "Off123456"}).json()["data"]
    login = client.post("/api/v1/auth/login", json={"username": "pytest_off", "password": "Off123456"}).json()
    headers = {"Authorization": f"Bearer {login['data']['access_token']}"}
    assert client.get("/api/v1/auth/me", headers=headers).status_code == 200
    client.put(f"/api/v1/accounts/{created['id']}", headers=admin,
               json={"full_name": "Sắp khóa", "role": "CTSV", "active": False})
    assert client.get("/api/v1/auth/me", headers=headers).status_code == 401
    assert client.post("/api/v1/auth/login", json={"username": "pytest_off", "password": "Off123456"}).status_code == 403


# ---------------------------------------------------------------------------- areas/cameras

def test_camera_sensitive_fields_hidden_from_ctsv(client, admin, ctsv, setup_camera):
    admin_view = client.get("/api/v1/cameras", headers=admin).json()["items"]
    ctsv_view = client.get("/api/v1/cameras", headers=ctsv).json()["items"]
    cam_admin = next(c for c in admin_view if c["code"] == "CAM_A1")
    cam_ctsv = next(c for c in ctsv_view if c["code"] == "CAM_A1")
    assert cam_admin["stream_source"].startswith("rtsp://") and cam_admin["has_device_key"] is True
    assert "stream_source" not in cam_ctsv and "has_device_key" not in cam_ctsv
    assert all("KHOA_THIET_BI_HASH" not in c and "device_key" not in c for c in admin_view)
    areas = client.get("/api/v1/areas", headers=ctsv).json()["items"]
    assert {a["code"] for a in areas} >= {"CONG-A", "THU_VIEN"}
    dup = client.post("/api/v1/cameras", headers=admin, json={"code": "cam_a1", "name": "x",
                                                               "area_id": setup_camera["area_id"]})
    assert dup.status_code == 409


def test_device_authentication(client, setup_camera):
    ok = client.get("/api/v1/device/me", headers=setup_camera["device"])
    assert ok.status_code == 200 and ok.json()["data"]["area"]["name"] == "Cổng A"
    bad = client.get("/api/v1/device/me", headers={"X-Camera-Code": "CAM_A1", "X-Device-Key": "wrong"})
    assert bad.status_code == 401
    assert client.get("/api/v1/device/me").status_code == 401


# ------------------------------------------------------------------------------- ingestion

def test_ingest_creates_record_without_identity(client, admin, setup_camera, evidence_dir):
    response = send(client, setup_camera["device"], code="BOTH")
    assert response.status_code == 201, response.text
    record = response.json()["data"]
    assert record["type"]["code"] == "BOTH" and record["area"]["name"] == "Cổng A"
    assert record["evidence"]["status"] == "available"
    assert not any(key in record for key in ("student", "ma_sv", "ho_ten"))
    assert str(evidence_dir) not in response.text  # no internal path leaked
    image = client.get(record["evidence"]["url"], headers=admin)
    assert image.status_code == 200 and image.headers["content-type"] == "image/jpeg"
    assert image.content.startswith(b"\xff\xd8\xff")
    assert client.get(record["evidence"]["url"]).status_code == 401  # evidence needs a user session


def test_ingest_is_idempotent_on_event_id(client, database, setup_camera):
    event_id = str(uuid.uuid4())
    first = send(client, setup_camera["device"], event_id=event_id)
    second = send(client, setup_camera["device"], event_id=event_id)
    assert first.status_code == 201 and second.status_code == 200
    assert second.json()["duplicate"] is True and second.json()["data"]["id"] == first.json()["data"]["id"]
    count = database.fetch_one("SELECT COUNT(*) AS N FROM dbo.PHAT_HIEN_VIPHAM WHERE MA_SU_KIEN = ?", (event_id,))
    assert count["N"] == 1
    other_camera = send(client, setup_camera["device2"], event_id=event_id)
    assert other_camera.status_code == 409


@pytest.mark.parametrize("extra,status", [
    ({"ma_sv": "SV001"}, 422),
    ({"ho_ten": "Nguyễn Văn An"}, 422),
    ({"unexpected": "x"}, 422),
    ({"violation_code": "SMOKING"}, 422),
    ({"event_id": "not-a-uuid"}, 422),
    ({"uniform_confidence": "1.5"}, 422),
    ({"track_id": "-3"}, 422),
    ({"occurred_at": "13/10/2026"}, 422),
])
def test_ingest_rejects_invalid_metadata(client, setup_camera, extra, status):
    response = send(client, setup_camera["device"], extra=extra)
    assert response.status_code == status, response.text
    assert response.json()["error"]["code"] == "validation_error"


def test_ingest_rejects_bad_images_and_times(client, setup_camera):
    assert send(client, setup_camera["device"], image=b"GIF89a....").status_code == 415
    assert send(client, setup_camera["device"], image=b"\xff\xd8\xffbroken").status_code == 422
    future = local_now(client) + timedelta(hours=2)
    assert send(client, setup_camera["device"], occurred_at=future).status_code == 422
    old = local_now(client) - timedelta(days=30)
    assert send(client, setup_camera["device"], occurred_at=old).status_code == 422
    no_image = client.post("/api/v1/violations", headers=setup_camera["device"], data={
        "event_id": str(uuid.uuid4()), "violation_code": "UNIFORM", "occurred_at": local_now(client).isoformat()})
    assert no_image.status_code in (415, 422)
    big = client.post("/api/v1/violations", headers=setup_camera["device"],
                      data={"event_id": str(uuid.uuid4()), "violation_code": "UNIFORM",
                            "occurred_at": local_now(client).isoformat()},
                      files={"image": ("big.jpg", io.BytesIO(b"\xff\xd8\xff" + b"0" * (3 * 1024 * 1024)), "image/jpeg")})
    assert big.status_code == 413


def test_inactive_camera_and_inactive_violation_type_are_rejected(client, admin, database, setup_camera):
    camera_id = setup_camera["camera2_id"]
    payload = {"code": "CAM_TV1", "name": "Camera thư viện", "area_id": setup_camera["area2_id"], "active": False}
    assert client.put(f"/api/v1/cameras/{camera_id}", headers=admin, json=payload).status_code == 200
    assert send(client, setup_camera["device2"]).status_code == 403
    payload["active"] = True
    client.put(f"/api/v1/cameras/{camera_id}", headers=admin, json=payload)
    assert send(client, setup_camera["device2"]).status_code == 201

    with database.transaction() as cursor:
        cursor.execute("UPDATE dbo.LOI_VIPHAM SET DANG_SU_DUNG = 0 WHERE MA_LOI = 'NO_CARD'")
    try:
        response = send(client, setup_camera["device"], code="NO_CARD")
        assert response.status_code == 409 and response.json()["error"]["code"] == "violation_type_inactive"
    finally:
        with database.transaction() as cursor:
            cursor.execute("UPDATE dbo.LOI_VIPHAM SET DANG_SU_DUNG = 1 WHERE MA_LOI = 'NO_CARD'")


def test_rotated_device_key_invalidates_old_one(client, admin, setup_camera):
    old_headers = dict(setup_camera["device2"])
    rotated = client.post(f"/api/v1/cameras/{setup_camera['camera2_id']}/rotate-key", headers=admin)
    assert rotated.status_code == 200
    assert client.get("/api/v1/device/me", headers=old_headers).status_code == 401
    setup_camera["device2"]["X-Device-Key"] = rotated.json()["device_key"]
    assert client.get("/api/v1/device/me", headers=setup_camera["device2"]).status_code == 200


# ------------------------------------------------------------------------- lookup/processing

def test_list_filter_sort_and_paginate(client, ctsv, setup_camera):
    base = local_now(client).replace(hour=9, minute=0, second=0) - timedelta(days=1)
    for offset, code in enumerate(["UNIFORM", "NO_CARD", "UNIFORM"]):
        assert send(client, setup_camera["device2"], code=code,
                    occurred_at=base + timedelta(minutes=offset)).status_code == 201
    area_only = client.get("/api/v1/violations", headers=ctsv,
                           params={"area_id": setup_camera["area2_id"], "page_size": 2, "sort": "oldest"}).json()
    assert area_only["total"] >= 4 and len(area_only["items"]) == 2 and area_only["pages"] >= 2
    times = [item["occurred_at"] for item in area_only["items"]]
    assert times == sorted(times)
    types = client.get("/api/v1/violation-types", headers=ctsv).json()["items"]
    no_card_id = next(t["id"] for t in types if t["code"] == "NO_CARD")
    filtered = client.get("/api/v1/violations", headers=ctsv, params={
        "type_id": no_card_id, "date_from": base.date().isoformat(), "date_to": base.date().isoformat()}).json()
    assert filtered["total"] >= 1 and all(i["type"]["code"] == "NO_CARD" for i in filtered["items"])
    invalid = client.get("/api/v1/violations", headers=ctsv, params={"date_from": "2026-10-10", "date_to": "2026-10-01"})
    assert invalid.status_code == 422


def test_processing_update_never_assigns_identity(client, ctsv, setup_camera):
    record = send(client, setup_camera["device"]).json()["data"]
    response = client.patch(f"/api/v1/violations/{record['id']}", headers=ctsv,
                            json={"status": "DA_XU_LY", "note": "Đã nhắc nhở tại cổng"})
    assert response.status_code == 200
    processing = response.json()["data"]["processing"]
    assert processing["status"] == "DA_XU_LY" and processing["processed_at"] is not None
    assert processing["processed_by"] is None  # migration 007 dropped NGUOI_XU_LY_ID
    rejected = client.patch(f"/api/v1/violations/{record['id']}", headers=ctsv,
                            json={"status": "DA_XU_LY", "ma_sv": "SV001"})
    assert rejected.status_code == 422
    assert client.get("/api/v1/violations/999999999", headers=ctsv).status_code == 404


# ------------------------------------------------------------------------------- reports

def test_dashboard_and_reports(client, ctsv, admin, setup_camera):
    today = local_now(client).date()
    semester = client.post("/api/v1/semesters", headers=admin, json={
        "code": "HK1-2026", "name": "Học kỳ 1 năm học 2026-2027",
        "start_date": (today - timedelta(days=40)).isoformat(), "end_date": (today + timedelta(days=60)).isoformat()})
    assert semester.status_code == 201, semester.text
    overlap = client.post("/api/v1/semesters", headers=admin, json={
        "code": "HK-OVERLAP", "name": "Trùng", "start_date": today.isoformat(), "end_date": today.isoformat()})
    assert overlap.status_code == 409

    dashboard = client.get("/api/v1/dashboard", headers=ctsv, params={"period": "month"}).json()["data"]
    counters = dashboard["counters"]
    assert counters["today"] >= 1 and counters["week"] >= counters["today"] and counters["semester"] >= 1
    summary = dashboard["summary"]
    assert summary["total"] == sum(item["count"] for item in summary["by_type"])
    assert len(summary["by_hour"]) == 24 and summary["hotspots"]

    semester_report = client.get("/api/v1/reports/summary", headers=ctsv, params={
        "period": "semester", "semester_id": semester.json()["data"]["id"],
        "area_id": setup_camera["area_id"]}).json()["data"]
    assert all(area["id"] == setup_camera["area_id"] for area in semester_report["summary"]["by_area"])
    custom = client.get("/api/v1/reports/summary", headers=ctsv, params={
        "period": "custom", "date_from": (today - timedelta(days=2)).isoformat(), "date_to": today.isoformat()})
    assert custom.status_code == 200 and len(custom.json()["data"]["summary"]["by_day"]) == 3
    assert client.get("/api/v1/reports/summary", headers=ctsv, params={"period": "custom"}).status_code == 422


def test_report_exports(client, ctsv):
    from openpyxl import load_workbook

    xlsx = client.get("/api/v1/reports/export", headers=ctsv, params={"format": "xlsx", "period": "month"})
    assert xlsx.status_code == 200 and "attachment" in xlsx.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames[0] == "Tổng hợp" and "Chi tiết" in workbook.sheetnames
    assert workbook["Tổng hợp"]["A1"].value.startswith("BÁO CÁO THỐNG KÊ")
    assert workbook["Chi tiết"].max_row >= 2
    pdf = client.get("/api/v1/reports/export", headers=ctsv, params={"format": "pdf", "period": "week"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert client.get("/api/v1/reports/export", headers=ctsv, params={"format": "csv"}).status_code == 422


# ------------------------------------------------------------------------------- retention

def test_retention_deletes_files_but_keeps_statistics(client, admin, database, setup_camera, evidence_dir):
    from api.evidence import EvidenceStorage
    from api.retention import purge_expired_evidence

    records = [send(client, setup_camera["device"]).json()["data"] for _ in range(3)]
    ids = [r["id"] for r in records]
    keys = {row["PHAT_HIEN_ID"]: row["ANH_KHOA"] for row in database.fetch_all(
        "SELECT PHAT_HIEN_ID, ANH_KHOA FROM dbo.ANH_VIPHAM WHERE PHAT_HIEN_ID IN (?, ?, ?)", ids)}
    storage = EvidenceStorage(evidence_dir)
    storage.delete(keys[ids[1]])  # simulate a file removed manually before the job
    with database.transaction() as cursor:
        cursor.execute("UPDATE dbo.ANH_VIPHAM SET THOI_GIAN_HET_HAN = DATEADD(DAY, -1, SYSDATETIME()) "
                       "WHERE PHAT_HIEN_ID IN (?, ?)", ids[0], ids[1])
        cursor.execute("UPDATE dbo.ANH_VIPHAM SET THOI_GIAN_HET_HAN = DATEADD(DAY, -1, SYSDATETIME()), "
                       "ANH_KHOA = '../../outside.jpg' WHERE PHAT_HIEN_ID = ?", ids[2])
    outside = evidence_dir.parent / "outside.jpg"
    outside.write_bytes(b"must survive")
    total_before = database.fetch_one("SELECT COUNT(*) AS N FROM dbo.PHAT_HIEN_VIPHAM")["N"]

    report = purge_expired_evidence(database, storage, local_now(client))
    assert report.deleted >= 1 and report.already_missing >= 1 and report.unsafe >= 1 and report.errors == 0
    assert not storage.exists(keys[ids[0]])
    assert outside.exists()  # unsafe key never deletes outside the evidence root
    outside.unlink()

    total_after = database.fetch_one("SELECT COUNT(*) AS N FROM dbo.PHAT_HIEN_VIPHAM")["N"]
    assert total_after == total_before  # statistics rows are kept
    detail = client.get(f"/api/v1/violations/{ids[0]}", headers=admin).json()["data"]
    assert detail["evidence"]["status"] == "deleted" and detail["evidence"]["url"] is None
    assert detail["type"]["code"] and detail["area"]["name"] and detail["occurred_at"]
    gone = client.get(f"/api/v1/violations/{ids[0]}/evidence", headers=admin)
    assert gone.status_code == 410 and "xóa" in gone.json()["error"]["message"]
    cleared = database.fetch_all("SELECT ANH_KHOA FROM dbo.ANH_VIPHAM WHERE PHAT_HIEN_ID IN (?, ?, ?)", ids)
    assert len(cleared) == 3 and all(row["ANH_KHOA"] is None for row in cleared)  # image rows kept, keys cleared
    second_run = purge_expired_evidence(database, storage, local_now(client))
    assert second_run.deleted == 0 and second_run.errors == 0  # re-runnable, nothing left


def test_missing_file_is_reported_as_missing(client, admin, database, setup_camera, evidence_dir):
    from api.evidence import EvidenceStorage

    record = send(client, setup_camera["device"]).json()["data"]
    key = database.fetch_one("SELECT ANH_KHOA FROM dbo.ANH_VIPHAM WHERE PHAT_HIEN_ID = ?", (record["id"],))["ANH_KHOA"]
    EvidenceStorage(evidence_dir).delete(key)
    response = client.get(f"/api/v1/violations/{record['id']}/evidence", headers=admin)
    assert response.status_code == 410
    assert client.get(f"/api/v1/violations/{record['id']}", headers=admin).json()["data"]["evidence"]["status"] == "missing"


# --------------------------------------------------------------------- legacy table + failures

def test_ct_vipham_allows_repeated_violation_of_same_type(database):
    with database.transaction() as cursor:
        cursor.execute("INSERT INTO dbo.CT_VIPHAM (ID_LOI, ID_SV, NGAY_VIPHAM, GHI_CHU) OUTPUT INSERTED.ID "
                       "VALUES (1, 1, '2026-10-12', N'pytest lần 1')")
        first = cursor.fetchone()[0]
        cursor.execute("INSERT INTO dbo.CT_VIPHAM (ID_LOI, ID_SV, NGAY_VIPHAM, GHI_CHU) OUTPUT INSERTED.ID "
                       "VALUES (1, 1, '2026-10-13', N'pytest lần 2')")
        second = cursor.fetchone()[0]
    try:
        assert first != second
        rows = database.fetch_all("SELECT ID FROM dbo.CT_VIPHAM WHERE ID_LOI = 1 AND ID_SV = 1")
        assert len(rows) >= 3  # sample row + two new ones
    finally:
        with database.transaction() as cursor:
            cursor.execute("DELETE FROM dbo.CT_VIPHAM WHERE ID IN (?, ?)", first, second)


def test_database_unavailable_returns_503(client, admin):
    from dataclasses import replace

    from api.database import Database
    from api.dependencies import get_db

    settings = client.app.state.settings.database
    broken = Database(replace(settings, server="tcp:127.0.0.1,1", timeout=2))
    client.app.dependency_overrides[get_db] = lambda: broken
    try:
        response = client.get("/api/v1/areas", headers=admin)
        assert response.status_code == 503 and response.json()["error"]["code"] == "database_unavailable"
        health = client.get("/api/v1/health")
        assert health.status_code == 503 and health.json()["database"]["status"] == "unavailable"
        assert "127.0.0.1" not in health.text
    finally:
        client.app.dependency_overrides.clear()
